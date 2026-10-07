#!/usr/bin/env python3
"""Read-only component and global verification for the production migration."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import openpyxl

from production_sheet_gate import (
    OVERLAY_HEADERS,
    PACKAGE_PATHS,
    PROTECTED_SHEETS,
    TRAINER_METADATA_SHEETS,
    a1_column_to_index,
    compare_fields,
    learning_headers,
    load_json,
    normalize_actual_for_expected,
    quiz_headers,
    sheet_hash,
    trainer_row,
)


def verify_common(workbook: Any, baseline: dict[str, Any], errors: list[dict[str, Any]]) -> None:
    for name in PROTECTED_SHEETS:
        current = sheet_hash(workbook[name])
        expected = baseline["protectedSheets"][name]
        if current != expected:
            errors.append({"component": "protected", "sheet": name, "expected": expected, "actual": current})
    for name in TRAINER_METADATA_SHEETS:
        current = sheet_hash(workbook[name])
        expected = baseline["trainerMetadata"][name]
        if current != expected:
            errors.append({"component": "trainer-metadata", "sheet": name, "expected": expected, "actual": current})


def verify_trainer(workbook: Any, package: dict[str, Any], errors: list[dict[str, Any]]) -> dict[str, Any]:
    passed = 0
    for item in package["records"]:
        ws = workbook[item["sourceSheet"]]
        try:
            row = trainer_row(ws, item["id"])
        except ValueError as error:
            errors.append({"component": "trainer", "id": item["id"], "error": str(error)})
            continue
        final = dict(item["expectedFields"])
        final.update(item["changes"])
        mismatches = []
        for column, wanted in final.items():
            cell = ws.cell(row=row, column=a1_column_to_index(column) + 1)
            actual = normalize_actual_for_expected(cell.value, wanted, cell.number_format)
            if actual != wanted:
                mismatches.append({"field": column, "expected": wanted, "actual": actual})
        if mismatches:
            errors.append({"component": "trainer", "id": item["id"], "row": row, "mismatches": mismatches})
        else:
            passed += 1
    return {"ids": passed, "expectedIds": 9, "cells": sum(len(item["writeColumns"]) for item in package["records"])}


def verify_quiz(workbook: Any, package: dict[str, Any], errors: list[dict[str, Any]]) -> dict[str, Any]:
    ws = workbook["Quizfragen"]
    headers = quiz_headers(ws)
    key_rows: dict[str, list[int]] = {}
    id_rows: dict[str, list[int]] = {}
    for row in range(2, ws.max_row + 1):
        key = str(ws.cell(row, 1).value or "").strip()
        identifier = str(ws.cell(row, 3).value or "").strip()
        if key:
            key_rows.setdefault(key, []).append(row)
        if identifier:
            id_rows.setdefault(identifier, []).append(row)
    creates = 0
    updates = 0
    for item in package["changes"]:
        matches = sorted(set(key_rows.get(item["quizKey"], []) + id_rows.get(item["id"], [])))
        if len(matches) != 1:
            errors.append({"component": "quiz", "id": item["id"], "error": f"erwartet 1 Zeile, gefunden {len(matches)}"})
            continue
        row = matches[0]
        mismatches = compare_fields(ws, row, headers, item["expectedAfterFields"])
        if mismatches:
            errors.append({"component": "quiz", "id": item["id"], "row": row, "mismatches": mismatches})
            continue
        if item["action"] == "create":
            creates += 1
        else:
            updates += 1
    active = sum(1 for row in range(2, ws.max_row + 1) if str(ws.cell(row, 9).value or "").strip().lower() == "ja")
    return {"creates": creates, "updates": updates, "active": active, "uniqueKeys": len(key_rows), "uniqueIds": len(id_rows)}


def expected_overlay_rows(package: dict[str, Any]) -> list[list[str]]:
    rows = [list(OVERLAY_HEADERS)]
    for item in package["changes"]:
        fields = item.get("fieldsToWrite") or {}
        rows.append([
            item["id"], item["source"], "NEU" if item["kind"] == "create" else "OVERRIDE", item["topic"],
            fields.get("front", ""), fields.get("back", ""), "ja", package["revision"],
        ])
    return rows


def verify_cards(workbook: Any, package: dict[str, Any], errors: list[dict[str, Any]]) -> dict[str, Any]:
    if "Karteikarten_Ergaenzungen" not in workbook.sheetnames:
        errors.append({"component": "cards", "error": "Overlaytabelle fehlt"})
        return {"rows": 0, "new": 0, "overrides": 0}
    ws = workbook["Karteikarten_Ergaenzungen"]
    expected = expected_overlay_rows(package)
    actual = [[str(ws.cell(row, col).value or "") for col in range(1, 9)] for row in range(1, len(expected) + 1)]
    if actual != expected:
        errors.append({"component": "cards", "error": "Overlayinhalt weicht vom freigegebenen Paket ab"})
    ids = [row[0] for row in actual[1:]]
    if len(ids) != len(set(ids)):
        errors.append({"component": "cards", "error": "Doppelte Karten-ID im Overlay"})
    return {
        "rows": len(actual) - 1,
        "new": sum(1 for row in actual[1:] if row[2] == "NEU"),
        "overrides": sum(1 for row in actual[1:] if row[2] == "OVERRIDE"),
        "projectedActive": 2686 + sum(1 for row in actual[1:] if row[2] == "NEU"),
    }


def verify_learning(workbook: Any, package: dict[str, Any], errors: list[dict[str, Any]]) -> dict[str, Any]:
    ws = workbook["Lerntexte"]
    headers = learning_headers(ws)
    rows: dict[str, list[int]] = {}
    for row in range(3, ws.max_row + 1):
        identifier = str(ws.cell(row, 1).value or "").strip()
        if identifier:
            rows.setdefault(identifier, []).append(row)
    passed = 0
    for item in package["records"]:
        matches = rows.get(item["id"], [])
        if len(matches) != 1:
            errors.append({"component": "learning", "id": item["id"], "error": f"erwartet 1 Zeile, gefunden {len(matches)}"})
            continue
        mismatches = compare_fields(ws, matches[0], headers, item["finalValues"])
        if mismatches:
            errors.append({"component": "learning", "id": item["id"], "row": matches[0], "mismatches": mismatches})
        else:
            passed += 1
    active = sum(1 for row in range(3, ws.max_row + 1) if str(ws.cell(row, 13).value or "").strip().lower() == "ja")
    chapters = set()
    for row in range(3, ws.max_row + 1):
        if str(ws.cell(row, 13).value or "").strip().lower() == "ja":
            chapters.add((str(ws.cell(row, 2).value or ""), str(ws.cell(row, 4).value or "")))
    return {"ids": passed, "active": active, "chapters": len(chapters)}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--baseline", required=True, type=Path)
    parser.add_argument("--components", required=True, help="comma-separated: trainer,quiz,cards,learning")
    parser.add_argument("--report", required=True, type=Path)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.xlsx, data_only=False, read_only=False)
    baseline = load_json(args.baseline)
    packages = {name: load_json(path) for name, path in PACKAGE_PATHS.items()}
    errors: list[dict[str, Any]] = []
    verify_common(workbook, baseline, errors)
    results: dict[str, Any] = {}
    for component in [item.strip() for item in args.components.split(",") if item.strip()]:
        if component == "trainer":
            results[component] = verify_trainer(workbook, packages["trainer"], errors)
        elif component == "quiz":
            results[component] = verify_quiz(workbook, packages["quiz"], errors)
        elif component == "cards":
            results[component] = verify_cards(workbook, packages["cards"], errors)
        elif component == "learning":
            results[component] = verify_learning(workbook, packages["learning"], errors)
        else:
            errors.append({"component": component, "error": "Unbekannte Komponente"})
    report = {
        "schemaVersion": 1,
        "mode": "READ_ONLY_POSTWRITE_VERIFICATION",
        "status": "PASS" if not errors else "FAIL",
        "components": results,
        "protectedSheetsUnchanged": not any(error["component"] == "protected" for error in errors),
        "trainerMetadataUnchanged": not any(error["component"] == "trainer-metadata" for error in errors),
        "errors": errors,
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    return 0 if report["status"] == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
